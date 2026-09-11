from datetime import datetime
from io import BytesIO

from docx import Document

from .base import Exporter
from ..models import TranscriptionResult, TranscriptionStatus


class DocxExporter(Exporter):
    extension = "docx"

    def render_individual(self, result: TranscriptionResult) -> bytes:
        doc = Document()
        doc.add_heading(result.audio_file.path.name, level=1)

        meta = doc.add_paragraph()
        run = meta.add_run(
            f"Duração: {result.duration_str}  |  Idioma: {result.language}  |  Confiança: {result.confidence_pct}"
        )
        run.italic = True

        doc.add_paragraph()
        if result.status == TranscriptionStatus.COMPLETED:
            for line in self._body_lines(result):
                doc.add_paragraph(line)
        else:
            doc.add_paragraph(f"[FALHOU: {result.error}]")

        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def render_full(self, results: list[TranscriptionResult]) -> bytes:
        now = datetime.now().strftime("%Y-%m-%d %H:%M")
        folder = results[0].audio_file.path.parent.name if results else "—"
        lang = results[0].language if results else "—"

        doc = Document()
        doc.add_heading(f"Transcrição — {folder}", level=1)

        meta = doc.add_paragraph()
        run = meta.add_run(f"Gerado em {now}  ·  {len(results)} áudios  ·  Idioma: {lang}")
        run.italic = True

        doc.add_heading("Índice", level=2)
        for r in results:
            doc.add_paragraph(f"{r.id}. {r.audio_file.path.name} ({r.duration_str})")

        doc.add_page_break()

        for r in results:
            doc.add_heading(f"{r.id} · {r.audio_file.path.name}", level=2)

            dur_p = doc.add_paragraph()
            dur_p.add_run(f"Duração: {r.duration_str}").italic = True

            if r.status == TranscriptionStatus.COMPLETED:
                for line in self._body_lines(r):
                    doc.add_paragraph(line)
            else:
                doc.add_paragraph(f"[FALHOU: {r.error}]")
            doc.add_paragraph("─" * 30)

        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()
