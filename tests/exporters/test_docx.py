from io import BytesIO
from docx import Document
from src.exporters.docx import DocxExporter


def test_docx_individual_returns_bytes(one_result):
    exporter = DocxExporter()
    output = exporter.render_individual(one_result)
    assert isinstance(output, bytes)
    assert len(output) > 0


def test_docx_individual_is_valid_docx(one_result):
    exporter = DocxExporter()
    output = exporter.render_individual(one_result)
    doc = Document(BytesIO(output))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Olá, como vai você?" in full_text


def test_docx_individual_has_filename_heading(one_result):
    exporter = DocxExporter()
    output = exporter.render_individual(one_result)
    doc = Document(BytesIO(output))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any("WhatsApp Audio" in h for h in headings)


def test_docx_full_returns_bytes(two_results):
    exporter = DocxExporter()
    output = exporter.render_full(two_results)
    assert isinstance(output, bytes)


def test_docx_full_contains_all_texts(two_results):
    exporter = DocxExporter()
    output = exporter.render_full(two_results)
    doc = Document(BytesIO(output))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Primeiro áudio." in full_text
    assert "Segundo áudio." in full_text


def test_export_writes_docx_files(tmp_path, two_results):
    exporter = DocxExporter()
    written = exporter.export(two_results, tmp_path)
    assert any(p.suffix == ".docx" and "FULL" in p.name for p in written)
    assert all(p.exists() for p in written)
